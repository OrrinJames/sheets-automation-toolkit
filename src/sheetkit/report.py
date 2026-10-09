"""Formatted Excel summary reports built with pandas + openpyxl."""

from __future__ import annotations

from pathlib import Path
from typing import Literal

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from sheetkit.clean import normalise_header

AggFunc = Literal["sum", "mean", "count", "min", "max"]

HEADER_FILL = PatternFill("solid", start_color="1F4E78", end_color="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FONT = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NZD_FORMAT = '"$"#,##0.00;[Red]-"$"#,##0.00'


def build_summary(
    df: pd.DataFrame,
    index: str,
    values: str,
    columns: str | None = None,
    aggfunc: AggFunc = "sum",
) -> pd.DataFrame:
    """Build a pivot table with a ``Total`` row (and column, when pivoting by ``columns``).

    Args:
        df: Cleaned input table. ``values`` should be numeric.
        index: Column whose values become the rows.
        values: Numeric column to aggregate.
        columns: Optional column whose values become the pivot columns.
        aggfunc: Aggregation to apply.

    Returns:
        The pivot as a flat DataFrame, ``index`` as its first column.
    """
    index, values = normalise_header(index), normalise_header(values)
    columns = normalise_header(columns) if columns else None
    for name in filter(None, (index, values, columns)):
        if name not in df.columns:
            raise KeyError(f"Column {name!r} not found. Available: {', '.join(df.columns)}")

    data = df.copy()
    if aggfunc != "count":
        data[values] = pd.to_numeric(data[values], errors="coerce")

    pivot = pd.pivot_table(
        data,
        index=index,
        columns=columns,
        values=values,
        aggfunc=aggfunc,
        fill_value=0,
        margins=True,
        margins_name="Total",
        observed=True,
    )
    if isinstance(pivot, pd.Series):  # pragma: no cover - defensive, pandas returns a frame
        pivot = pivot.to_frame()
    if columns is None:
        pivot.columns = [values]
    else:
        pivot.columns = [str(c) for c in pivot.columns]
    pivot = pivot.reset_index()
    pivot[index] = pivot[index].astype(str)
    return pivot


def _style_header(ws: Worksheet, row: int, ncols: int) -> None:
    """Apply the header fill, font and border to one row."""
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def autofit_columns(ws: Worksheet, min_width: int = 8, max_width: int = 60) -> None:
    """Size each column to its longest rendered value, within sensible bounds."""
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            value = cell.value
            if isinstance(value, float):
                text = f"{value:,.2f}"
            elif hasattr(value, "year"):
                text = "dd/mm/yyyy"
            else:
                text = str(value)
            widths[cell.column] = max(widths.get(cell.column, 0), len(text))
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(
            min_width, min(max_width, width + 2)
        )


def _write_frame(
    ws: Worksheet,
    df: pd.DataFrame,
    money: bool = False,
    total_row: bool = False,
) -> None:
    """Write a DataFrame to ``ws`` with a styled header, borders and frozen header row."""
    ws.append([str(c) for c in df.columns])
    _style_header(ws, 1, len(df.columns))
    for record in df.itertuples(index=False):
        ws.append([None if pd.isna(v) else _to_cell(v) for v in record])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = BORDER
            if money and isinstance(cell.value, int | float) and cell.column > 1:
                cell.number_format = NZD_FORMAT
            elif hasattr(cell.value, "year"):
                cell.number_format = "dd/mm/yyyy"
    if total_row and ws.max_row > 1:
        for cell in ws[ws.max_row]:
            cell.font = TOTAL_FONT
    ws.freeze_panes = "A2"
    autofit_columns(ws)


def _to_cell(value: object) -> object:
    """Convert numpy/pandas scalars into types openpyxl writes natively."""
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item"):
        return value.item()
    return value


def write_report(
    df: pd.DataFrame,
    output: str | Path,
    index: str,
    values: str,
    columns: str | None = None,
    aggfunc: AggFunc = "sum",
    title: str = "Summary",
    include_data: bool = True,
) -> Path:
    """Write a formatted Excel workbook: a pivot summary sheet plus the source data.

    The summary sheet has a coloured, frozen header row, NZD number formatting
    (except for ``count``), a bold total row, and auto-fitted column widths.

    Args:
        df: Cleaned input table.
        output: Destination ``.xlsx`` path.
        index: Pivot row column.
        values: Column to aggregate.
        columns: Optional pivot column.
        aggfunc: Aggregation to apply.
        title: Name of the summary worksheet (max 31 characters).
        include_data: Also write the full input table to a ``Data`` sheet.

    Returns:
        The path of the written workbook.
    """
    output = Path(output)
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Report output must be an .xlsx file")
    summary = build_summary(df, index=index, values=values, columns=columns, aggfunc=aggfunc)

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    _write_frame(ws, summary, money=aggfunc != "count", total_row=True)

    if include_data:
        _write_frame(wb.create_sheet("Data"), df)

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output
