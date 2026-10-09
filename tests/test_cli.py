from __future__ import annotations

from openpyxl import load_workbook
from typer.testing import CliRunner

from sheetkit.cli import app
from sheetkit.io import read_table

runner = CliRunner()


def test_clean_command(fixtures_dir, tmp_path) -> None:
    out = tmp_path / "clean.csv"
    result = runner.invoke(
        app,
        [
            "clean",
            str(fixtures_dir / "messy.csv"),
            "-o",
            str(out),
            "-d",
            "order date",
            "-c",
            "Amount (NZD)",
        ],
    )
    assert result.exit_code == 0, result.output
    df = read_table(out)
    assert df["amount_nzd"].tolist() == ["1250.0", "89.9", "-20.0", "410.0"]
    assert df["order_date"].iloc[0] == "2026-04-03"


def test_dedupe_command(fixtures_dir, tmp_path) -> None:
    out = tmp_path / "unique.csv"
    result = runner.invoke(
        app, ["dedupe", str(fixtures_dir / "dupes.csv"), "-o", str(out), "-k", "email"]
    )
    assert result.exit_code == 0, result.output
    assert "Removed 2 duplicate" in result.output
    assert len(read_table(out)) == 2


def test_merge_command(fixtures_dir, tmp_path) -> None:
    out = tmp_path / "merged.xlsx"
    result = runner.invoke(
        app,
        [
            "merge",
            str(fixtures_dir / "part_a.csv"),
            str(fixtures_dir / "part_b.csv"),
            "-o",
            str(out),
        ],
    )
    assert result.exit_code == 0, result.output
    assert len(read_table(out)) == 3


def test_report_command(fixtures_dir, tmp_path) -> None:
    out = tmp_path / "report.xlsx"
    result = runner.invoke(
        app,
        [
            "report",
            str(fixtures_dir / "messy.csv"),
            "-o",
            str(out),
            "-i",
            "region",
            "-v",
            "amount_nzd",
            "-c",
            "amount_nzd",
        ],
    )
    assert result.exit_code == 0, result.output
    ws = load_workbook(out)["Summary"]
    rows = {r[0]: r[1] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows == {"Auckland": 1230.0, "Christchurch": 410.0, "Wellington": 89.9, "Total": 1729.9}


def test_report_bad_agg(fixtures_dir, tmp_path) -> None:
    result = runner.invoke(
        app,
        [
            "report",
            str(fixtures_dir / "messy.csv"),
            "-o",
            str(tmp_path / "r.xlsx"),
            "-i",
            "region",
            "-v",
            "amount_nzd",
            "--agg",
            "median",
        ],
    )
    assert result.exit_code != 0


def test_missing_column_gives_clean_error(fixtures_dir, tmp_path) -> None:
    result = runner.invoke(
        app,
        ["dedupe", str(fixtures_dir / "dupes.csv"), "-o", str(tmp_path / "x.csv"), "-k", "phone"],
    )
    assert result.exit_code == 2
