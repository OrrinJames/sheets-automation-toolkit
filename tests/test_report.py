from __future__ import annotations

import pandas as pd
import pytest
from openpyxl import load_workbook

from sheetkit.report import HEADER_FILL, build_summary, write_report


@pytest.fixture
def sales() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "region": ["Auckland", "Auckland", "Wellington", "Christchurch"],
            "product": ["Beans", "Milk", "Beans", "Beans"],
            "amount": [100.0, 50.0, 25.5, 10.0],
        }
    )


def test_build_summary_simple(sales: pd.DataFrame) -> None:
    out = build_summary(sales, index="region", values="amount")
    totals = dict(zip(out["region"], out["amount"], strict=True))
    assert totals == {"Auckland": 150.0, "Christchurch": 10.0, "Wellington": 25.5, "Total": 185.5}


def test_build_summary_with_columns(sales: pd.DataFrame) -> None:
    out = build_summary(sales, index="Region", values="Amount", columns="product")
    assert list(out.columns) == ["region", "Beans", "Milk", "Total"]
    total_row = out[out["region"] == "Total"].iloc[0]
    assert total_row["Beans"] == 135.5
    assert total_row["Total"] == 185.5


def test_build_summary_count(sales: pd.DataFrame) -> None:
    out = build_summary(sales, index="region", values="amount", aggfunc="count")
    assert out.set_index("region").loc["Auckland", "amount"] == 2


def test_write_report_formatting(sales: pd.DataFrame, tmp_path) -> None:
    path = write_report(sales, tmp_path / "report.xlsx", index="region", values="amount")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Data"]

    ws = wb["Summary"]
    assert ws["A1"].value == "region"
    assert ws["A1"].font.bold
    assert ws["A1"].fill.start_color.rgb.endswith(HEADER_FILL.start_color.rgb[-6:])
    assert ws.freeze_panes == "A2"
    assert ws.cell(row=ws.max_row, column=1).value == "Total"
    assert ws.cell(row=ws.max_row, column=1).font.bold
    assert ws["B2"].number_format.startswith('"$"')
    assert ws.column_dimensions["A"].width >= len("Christchurch")


def test_write_report_rejects_non_xlsx(sales: pd.DataFrame, tmp_path) -> None:
    with pytest.raises(ValueError):
        write_report(sales, tmp_path / "report.csv", index="region", values="amount")


def test_build_summary_missing_column(sales: pd.DataFrame) -> None:
    with pytest.raises(KeyError):
        build_summary(sales, index="store", values="amount")
